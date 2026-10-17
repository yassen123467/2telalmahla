from fastapi import FastAPI, File, UploadFile, HTTPException, Form
from fastapi.responses import JSONResponse, StreamingResponse
import uvicorn
import os
import shutil
from pathlib import Path
import fitz  # PyMuPDF
import camelot
import pandas as pd
from docx import Document
from openpyxl import load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from reportlab.lib.units import inch
import io
from PIL import Image

app = FastAPI(title="SmartConvert AI", description="Intelligent File Conversion API")

# Directory for temporary files
UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)

@app.post("/uploadfile/")
async def create_upload_file(file: UploadFile = File(...)):
    """
    Upload a file and return basic information.
    """
    # Validate file
    if not file.filename:
        raise HTTPException(status_code=400, detail="No file uploaded")
    
    # Save the file temporarily
    file_path = UPLOAD_DIR / file.filename
    try:
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
    finally:
        file.file.close()
    
    return JSONResponse(
        content={
            "filename": file.filename,
            "content_type": file.content_type,
            "size": file_path.stat().st_size,
            "message": "File uploaded successfully"
        }
    )

@app.post("/convert/")
async def convert_file(
    file: UploadFile = File(...),
    target_format: str = Form(...)
):
    """
    Convert an uploaded file to the target format.
    Supported conversions:
    - PDF to text, CSV, DOCX
    - Text to PDF
    - Image to PDF
    - Excel to PDF (basic)
    - Word to PDF (basic)
    """
    # Validate file
    if not file.filename:
        raise HTTPException(status_code=400, detail="No file uploaded")
    
    # Save the file temporarily
    file_path = UPLOAD_DIR / file.filename
    try:
        with open(file_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
    finally:
        file.file.close()
    
    # Determine the conversion based on file extension and target format
    file_ext = file_path.suffix.lower()
    target_format = target_format.lower()
    
    # PDF conversions
    if file_ext == ".pdf":
        if target_format == "txt":
            return await pdf_to_txt(file_path)
        elif target_format == "csv":
            return await pdf_to_csv(file_path)
        elif target_format == "docx":
            return await pdf_to_docx(file_path)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported target format {target_format} for PDF")
    
    # Text conversions
    elif file_ext == ".txt":
        if target_format == "pdf":
            return await txt_to_pdf(file_path)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported target format {target_format} for TXT")
    
    # Image conversions
    elif file_ext in [".png", ".jpg", ".jpeg", ".bmp", ".tiff"]:
        if target_format == "pdf":
            return await image_to_pdf(file_path)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported target format {target_format} for image")
    
    # Excel conversions
    elif file_ext in [".xlsx", ".xls"]:
        if target_format == "pdf":
            return await excel_to_pdf(file_path)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported target format {target_format} for Excel")
    
    # Word conversions
    elif file_ext == ".docx":
        if target_format == "pdf":
            return await docx_to_pdf(file_path)
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported target format {target_format} for Word")
    
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported file format {file_ext}")

async def pdf_to_txt(file_path: Path):
    """Extract text from PDF and return as a text file."""
    text = ""
    doc = fitz.open(file_path)
    for page in doc:
        text += page.get_text()
    doc.close()
    
    # Create a text file in memory
    output = io.BytesIO()
    output.write(text.encode('utf-8'))
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="text/plain",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.txt"}
    )

async def pdf_to_csv(file_path: Path):
    """Extract tables from PDF and return as CSV."""
    try:
        # Try to extract tables with camelot
        tables = camelot.read_pdf(str(file_path), pages='all')
        if tables.n == 0:
            # Fallback to text extraction and create a single column CSV
            text = ""
            doc = fitz.open(file_path)
            for page in doc:
                text += page.get_text()
            doc.close()
            # Create a CSV with one column
            df = pd.DataFrame({'text': text.split('\n')})
        else:
            # Combine all tables into one DataFrame
            df = tables[0].df
            for i in range(1, tables.n):
                df = df.append(tables[i].df, ignore_index=True)
    except Exception as e:
        # If camelot fails, fall back to text
        text = ""
        doc = fitz.open(file_path)
        for page in doc:
            text += page.get_text()
        doc.close()
        df = pd.DataFrame({'text': text.split('\n')})
    
    # Create CSV in memory
    output = io.BytesIO()
    df.to_csv(output, index=False)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.csv"}
    )

async def pdf_to_docx(file_path: Path):
    """Convert PDF to Word document (basic text extraction)."""
    doc = Document()
    pdf_doc = fitz.open(file_path)
    for page in pdf_doc:
        text = page.get_text()
        doc.add_paragraph(text)
    pdf_doc.close()
    
    # Save to a BytesIO object
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.docx"}
    )

async def txt_to_pdf(file_path: Path):
    """Convert text file to PDF."""
    # Read the text file
    with open(file_path, 'r', encoding='utf-8') as f:
        text = f.read()
    
    # Create a PDF in memory
    output = io.BytesIO()
    pdf = SimpleDocTemplate(output, pagesize=letter)
    styles = getSampleStyleSheet()
    style = styles["Normal"]
    
    # Split text into paragraphs
    paragraphs = text.split('\n')
    flowables = []
    for para in paragraphs:
        if para.strip():
            flowables.append(Paragraph(para, style))
    
    pdf.build(flowables)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.pdf"}
    )

async def image_to_pdf(file_path: Path):
    """Convert image to PDF."""
    # Open image
    img = Image.open(file_path)
    # Convert to RGB if necessary (for saving as PDF)
    if img.mode != 'RGB':
        img = img.convert('RGB')
    
    # Save to a BytesIO object as PDF
    output = io.BytesIO()
    img.save(output, format='PDF')
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.pdf"}
    )

async def excel_to_pdf(file_path: Path):
    """Convert Excel to PDF (basic)."""
    # Load the workbook
    wb = load_workbook(file_path)
    # We'll take the first sheet for simplicity
    ws = wb.active
    
    # Create a PDF in memory
    output = io.BytesIO()
    pdf = SimpleDocTemplate(output, pagesize=letter)
    styles = getSampleStyleSheet()
    style = styles["Normal"]
    
    # Extract data from the worksheet
    data = []
    for row in ws.iter_rows(values_only=True):
        data.append([str(cell) if cell is not None else '' for cell in row])
    
    # Create a table
    if data:
        t = Table(data)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 14),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        flowables = [t]
    else:
        flowables = [Paragraph("No data found in the Excel file.", style)]
    
    pdf.build(flowables)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.pdf"}
    )

async def docx_to_pdf(file_path: Path):
    """Convert Word to PDF (basic)."""
    # Load the document
    doc = Document(file_path)
    
    # Create a PDF in memory
    output = io.BytesIO()
    pdf = SimpleDocTemplate(output, pagesize=letter)
    styles = getSampleStyleSheet()
    style = styles["Normal"]
    
    # Extract text from the document
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    
    text = '\n'.join(full_text)
    paragraphs = text.split('\n')
    flowables = []
    for para in paragraphs:
        if para.strip():
            flowables.append(Paragraph(para, style))
    
    pdf.build(flowables)
    output.seek(0)
    
    return StreamingResponse(
        output,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={file_path.stem}.pdf"}
    )

@app.get("/")
async def root():
    return {"message": "SmartConvert AI API is running"}

if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=8000)